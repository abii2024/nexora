"""Genereer docs/logboek/logboek.docx — examen-logboek PAL."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "docs" / "logboek" / "logboek.docx"

HEADER_BG = "4A3358"
BODY_BG = "EDF4EF"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
FONT = "Calibri"


def shade(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def set_borders(table, color: str = "FFFFFF") -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), color)
        borders.append(b)
    tblPr.append(borders)


def set_cell_margins(cell, top=60, bottom=60, left=120, right=120) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def write_cell(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: RGBColor = BLACK,
    size: int = 10,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def set_row_height(row, mm: float, exact: bool = False) -> None:
    row.height = Mm(mm)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exact else WD_ROW_HEIGHT_RULE.AT_LEAST


def add_persoonsinfo(doc: Document) -> None:
    rows = [
        ("Naam kandidaat", "Abdisamad Guled Abdulle"),
        ("Studentnummer", "9022236"),
        ("Klas/groep", "PALBVSOD3A"),
        ("(Leer)bedrijf", "100812537"),
        ("Beoordelaar 1", "Ömer Eryigit"),
        ("Beoordelaar 2", ""),
    ]
    table = doc.add_table(rows=len(rows) + 1, cols=2)
    table.autofit = False
    set_borders(table)

    header = table.rows[0].cells
    header[0].merge(header[1])
    write_cell(
        header[0],
        "Persoonsinformatie",
        bold=True,
        color=WHITE,
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    shade(header[0], HEADER_BG)
    set_row_height(table.rows[0], 7)

    for i, (label, value) in enumerate(rows, start=1):
        write_cell(table.rows[i].cells[0], label, bold=True)
        write_cell(table.rows[i].cells[1], value)
        shade(table.rows[i].cells[0], BODY_BG)
        shade(table.rows[i].cells[1], BODY_BG)
        table.rows[i].cells[0].width = Cm(5)
        table.rows[i].cells[1].width = Cm(11)
        set_row_height(table.rows[i], 6)


def add_logboek(doc: Document) -> None:
    headers = ["Datum", "Werkzaamheden en/of examenopdracht", "Resultaat", "Werkproces"]
    rows = [
        (
            "14-03-2026",
            "Laravel 12 + GitHub-repo opgezet.",
            "Werkende dev-omgeving op main.",
            "W1, W3",
        ),
        (
            "21–22-04-2026",
            "DoD, ontwerpdocument, wireframes (32), ERD, use-case, flowchart, 16 user stories.",
            "Compleet ontwerp + backlog, PO-akkoord.",
            "W1, W2",
        ),
        (
            "23-04-2026",
            "Sprint 1: US-01 login, US-02 autorisatie, US-03 medewerker aanmaken, US-04 overzicht.",
            "70 tests groen, tag sprint-1.",
            "W3, W4",
        ),
        (
            "24-04-2026",
            "Sprint 2: US-05 bewerken, US-06 deactiveren, US-07 cliënt aanmaken, US-08 begeleiders koppelen.",
            "85 tests groen, tag sprint-2.",
            "W3, W4",
        ),
        (
            "24-04-2026",
            "Sprint 3: US-09 overzicht, US-10 archiveren, US-11 concept-uren, US-12 uren-workflow.",
            "117 tests groen, tag sprint-3.",
            "W3, W4",
        ),
        (
            "24-04-2026",
            "Sprint 4: US-13 goedkeuren, US-14 urenoverzicht, US-15 wachtwoord-reset, US-16 profiel + deployment.",
            "360 tests totaal, tag sprint-4, live op Railway.",
            "W3, W4, W5",
        ),
        (
            "08–13-05-2026",
            "Screenshots, US-17 Resend-mail, overleggen-doc, reflectie, examen-checklist, GitHub-bewijslast.",
            "Examen-dossier compleet.",
            "W5, K2-W1, K2-W2, K2-W3",
        ),
    ]

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.autofit = False
    set_borders(table)

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        write_cell(cell, h, bold=True, color=WHITE, size=11)
        shade(cell, HEADER_BG)
    set_row_height(table.rows[0], 8)

    widths = [Cm(2.4), Cm(8.0), Cm(4.2), Cm(2.4)]
    for r, row_data in enumerate(rows, start=1):
        for c, value in enumerate(row_data):
            cell = table.rows[r].cells[c]
            write_cell(cell, value)
            shade(cell, BODY_BG)
            cell.width = widths[c]
        set_row_height(table.rows[r], 10)


def add_handtekening(doc: Document) -> None:
    table = doc.add_table(rows=3, cols=4)
    table.autofit = False
    set_borders(table)

    header = table.rows[0].cells
    header[0].merge(header[1]).merge(header[2]).merge(header[3])
    write_cell(
        header[0],
        "Handtekening voor gezien",
        bold=True,
        color=WHITE,
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    shade(header[0], HEADER_BG)
    set_row_height(table.rows[0], 7)

    widths = [Cm(3.0), Cm(8.5), Cm(2.0), Cm(3.5)]
    for r, label in enumerate(("Beoordelaar", "Kandidaat"), start=1):
        cells = table.rows[r].cells
        write_cell(cells[0], label, bold=True)
        write_cell(cells[1], "")
        write_cell(cells[2], "Datum", bold=True)
        write_cell(cells[3], "")
        for c in range(4):
            shade(cells[c], BODY_BG)
            cells[c].width = widths[c]
        set_row_height(table.rows[r], 12, exact=True)


def add_spacer(doc: Document, pt: int = 4) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(pt)


def main() -> None:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10)

    for section in doc.sections:
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("Logboek kandidaat")
    run.bold = True
    run.font.name = FONT
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x4A, 0x33, 0x58)

    add_persoonsinfo(doc)
    add_spacer(doc)
    add_logboek(doc)
    add_spacer(doc)
    add_handtekening(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
